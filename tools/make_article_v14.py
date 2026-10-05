"""把 v1.4 更新公告排成可投 B 站专栏的 Word（封面 + 正片截图 + 分节排版）。

用法： python tools\\make_article_v14.py
产出： 公告_v1.4\\公告_v1.4_封面.png、公告_v1.4\\公告_v1.4_专栏版.docx
"""
from __future__ import annotations

import re
import sys
from pathlib import Path

from PIL import Image, ImageDraw, ImageFont

HERE = Path(__file__).resolve().parent.parent
MD = HERE / "更新公告_v1.4.md"
OUT_DIR = HERE / "公告_v1.4"
COVER = OUT_DIR / "公告_v1.4_封面.png"
DOCX = OUT_DIR / "公告_v1.4_专栏版.docx"
SHOT_GRAPH = HERE / ".fix12" / "layout_bounds.png"
SHOT_TREE = HERE / ".fix12" / "layout_measure.png"

FONTS = ["C:/Windows/Fonts/msyhbd.ttc", "C:/Windows/Fonts/msyh.ttc",
         "C:/Windows/Fonts/simhei.ttf"]


def font(size: int):
    for p in FONTS:
        if Path(p).exists():
            try:
                return ImageFont.truetype(p, size)
            except Exception:
                pass
    return ImageFont.load_default()


def make_cover() -> Path:
    W, H = 1600, 900
    img = Image.new("RGB", (W, H), "#12141a")
    d = ImageDraw.Draw(img)
    for y in range(H):                                   # 竖向渐变底色
        t = y / H
        d.line([(0, y), (W, y)], fill=(int(18 + 16 * t), int(20 + 20 * t), int(27 + 34 * t)))
    for i, cx, cy, r, col in ((0, 250, 210, 320, (58, 74, 140)), (1, 1380, 700, 380, (40, 96, 110))):
        ov = Image.new("RGB", (W, H), "#12141a")
        od = ImageDraw.Draw(ov)
        od.ellipse([cx - r, cy - r, cx + r, cy + r], fill=col)
        img = Image.blend(img, ov, 0.22)
    d = ImageDraw.Draw(img)
    d.text((110, 150), "图片标签工坊", font=font(96), fill="#ffffff")
    d.text((112, 262), "v1.4 更新公告", font=font(58), fill="#8fb6ff")
    d.text((112, 356), "本地离线 · 自动打标 / 审核 / 检索 / 训练", font=font(32), fill="#c6cede")
    d.text((112, 410), "分类边界 · 扇区布局 · Danbooru 词表汉化", font=font(28), fill="#8f9bb0")
    if SHOT_GRAPH.exists():                              # 拿真实界面截图当主视觉
        shot = Image.open(SHOT_GRAPH).convert("RGB")
        shot.thumbnail((700, 430))
        px, py = 840, 300
        d.rectangle([px - 6, py - 6, px + shot.width + 6, py + shot.height + 6], outline="#48506a", width=2)
        img.paste(shot, (px, py))
        d.text((112, 500), "新版图谱：分类=扇区、零重叠", font=font(26), fill="#7f8a9e")
        d.text((112, 540), "工具栏可开关「分类边界」", font=font(26), fill="#7f8a9e")
    d.text((112, 800), "适用：任意旧版本升级（更新包）/ 全新安装（完整安装包）", font=font(24), fill="#6c7688")
    OUT_DIR.mkdir(parents=True, exist_ok=True)
    img.save(COVER)
    return COVER


BOLD = re.compile(r"\*\*(.+?)\*\*")


def add_rich(par, text: str) -> None:
    """把 **加粗** 写成真的加粗 run，其余照抄。"""
    pos = 0
    for m in BOLD.finditer(text):
        if m.start() > pos:
            par.add_run(text[pos:m.start()])
        par.add_run(m.group(1)).bold = True
        pos = m.end()
    if pos < len(text):
        par.add_run(text[pos:])


def build_docx() -> Path:
    from docx import Document
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.shared import Inches, Pt, RGBColor
    from docx.oxml.ns import qn

    doc = Document()
    style = doc.styles["Normal"]
    style.font.name = "微软雅黑"
    style.font.size = Pt(11)
    style.element.rPr.rFonts.set(qn("w:eastAsia"), "微软雅黑")
    for s in ("Heading 1", "Heading 2"):
        f = doc.styles[s].font
        f.name = "微软雅黑"
        f.color.rgb = RGBColor(0x1F, 0x3B, 0x73)
        f.size = Pt(17 if s == "Heading 1" else 14)
        doc.styles[s].element.rPr.rFonts.set(qn("w:eastAsia"), "微软雅黑")

    doc.add_picture(str(COVER), width=Inches(6.3))
    doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
    inserted = 0
    for raw in MD.read_text(encoding="utf-8").splitlines():
        line = raw.rstrip()
        if not line.strip():
            continue
        if line.startswith("# "):
            h = doc.add_heading(line[2:].strip(), level=1)
            h.alignment = WD_ALIGN_PARAGRAPH.CENTER
            continue
        if line.startswith("## "):
            doc.add_heading(line[3:].strip(), level=1)
            if inserted == 0 and SHOT_GRAPH.exists():       # 正文里插真实界面截图
                doc.add_picture(str(SHOT_GRAPH), width=Inches(6.3))
                doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
                doc.add_paragraph("新版图谱（分类=扇区、零重叠，可开关分类边界）").alignment = \
                    WD_ALIGN_PARAGRAPH.CENTER
                inserted += 1
            continue
        if line.startswith("- "):
            add_rich(doc.add_paragraph(style="List Bullet"), line[2:].strip())
            continue
        if re.match(r"^\d+\.\s", line):
            add_rich(doc.add_paragraph(style="List Number"), re.sub(r"^\d+\.\s", "", line))
            continue
        if line.startswith("|"):                            # 表格先跳过（md 表在专栏里不好看）
            continue
        add_rich(doc.add_paragraph(), line)
    doc.save(DOCX)
    return DOCX


def main() -> int:
    cover = make_cover()
    out = build_docx()
    print("封面:", cover)
    print("专栏版:", out, "%.1f MB" % (out.stat().st_size / 1048576))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
