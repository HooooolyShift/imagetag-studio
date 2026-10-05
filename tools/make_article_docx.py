r"""把 v1.3 更新公告排成可投 B 站专栏的 Word 文档（封面 + 配图 + 表格）。

用法： python tools\make_article_docx.py [输出.docx]
素材：公告_v1.3\公告_v1.3_封面.png、manual_shots\*.png（真实界面截图）
说明：正文按专栏阅读习惯排版（小标题 + 要点 + 配图），可直接复制进专栏编辑器。
"""
from __future__ import annotations

import sys
import tempfile
from pathlib import Path

from PIL import Image
from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

HERE = Path(__file__).resolve().parent.parent
SHOTS = HERE / "manual_shots"
OUT_DEFAULT = HERE / "公告_v1.3" / "公告_v1.3_专栏版.docx"

FONT = "Microsoft YaHei"
BLACK = RGBColor(0x00, 0x00, 0x00)
BODY_GREY = RGBColor(0x3A, 0x3F, 0x47)
CAP_GREY = RGBColor(0x6B, 0x72, 0x7E)
META_GREY = RGBColor(0x5B, 0x62, 0x70)
ACCENT = "2B6CB0"


def set_font(run, size: float | None = None, bold: bool | None = None,
             color: RGBColor | None = None, name: str = FONT) -> None:
    run.font.name = name
    run._element.rPr.rFonts.set(qn("w:ascii"), name)
    run._element.rPr.rFonts.set(qn("w:hAnsi"), name)
    run._element.rPr.rFonts.set(qn("w:eastAsia"), name)
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.font.bold = bold
    if color is not None:
        run.font.color.rgb = color


def para(doc, text: str = "", size: float = 11, bold: bool = False,
         color: RGBColor | None = BODY_GREY, space_after: float = 7,
         line_spacing: float = 1.35, align=None, indent: float = 0):
    p = doc.add_paragraph()
    if align is not None:
        p.alignment = align
    pf = p.paragraph_format
    pf.space_after = Pt(space_after)
    pf.space_before = Pt(0)
    pf.line_spacing = line_spacing
    if indent:
        pf.left_indent = Pt(indent)
    if text:
        set_font(p.add_run(text), size, bold, color)
    return p


def heading(doc, text: str, level: int = 1) -> None:
    p = doc.add_paragraph()
    pf = p.paragraph_format
    pf.space_before = Pt(16 if level == 1 else 12)
    pf.space_after = Pt(6 if level == 1 else 4)
    pf.line_spacing = 1.2
    set_font(p.add_run(text), 15.5 if level == 1 else 12.5, True, BLACK)
    p.style = doc.styles["Heading %d" % level]
    for run in p.runs:                      # 样式带颜色时压回纯黑
        run.font.color.rgb = BLACK


def bullet(doc, text: str, size: float = 11) -> None:
    p = doc.add_paragraph()
    pf = p.paragraph_format
    pf.left_indent = Pt(16)
    pf.first_line_indent = Pt(-11)
    pf.space_after = Pt(4)
    pf.line_spacing = 1.3
    set_font(p.add_run("•  "), size, False, RGBColor(0x2B, 0x6C, 0xB0))
    set_font(p.add_run(text), size, False, BODY_GREY)


def figure(doc, image: Path, caption: str, width_in: float = 5.0,
           keep_ratio: float | None = None) -> None:
    """插入配图。keep_ratio 用于裁掉界面截图下方的空白（只保留上面多少比例）。"""
    src = image
    if keep_ratio:
        im = Image.open(image)
        tmp = Path(tempfile.gettempdir()) / f"article_fig_{image.stem}_{int(keep_ratio*100)}.png"
        im.crop((0, 0, im.width, int(im.height * keep_ratio))).save(tmp)
        src = tmp
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(8)
    p.paragraph_format.space_after = Pt(3)
    p.add_run().add_picture(str(src), width=Inches(width_in))
    c = doc.add_paragraph()
    c.alignment = WD_ALIGN_PARAGRAPH.CENTER
    c.paragraph_format.space_after = Pt(12)
    set_font(c.add_run(caption), 9, False, CAP_GREY)


def shade(cell, hex_fill: str) -> None:
    tcPr = cell._tc.get_or_add_tcPr()
    el = OxmlElement("w:shd")
    el.set(qn("w:val"), "clear")
    el.set(qn("w:color"), "auto")
    el.set(qn("w:fill"), hex_fill)
    tcPr.append(el)


def table_borders(table, color: str = "D9D9D9") -> None:
    tbl = table._tbl
    tblPr = tbl.tblPr
    borders = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        el = OxmlElement(f"w:{edge}")
        el.set(qn("w:val"), "single")
        el.set(qn("w:sz"), "6")
        el.set(qn("w:color"), color)
        borders.append(el)
    tblPr.append(borders)


def keep_next(paragraph) -> None:
    """让这段和下一段留在同一页（标题不孤零零落在页尾）。"""
    pPr = paragraph._p.get_or_add_pPr()
    el = OxmlElement("w:keepNext")
    pPr.append(el)


def rows_keep_together(table) -> None:
    """整张表不断行：行内不拆 + 除最后一行外都与下一行同页。"""
    rows = table.rows
    for idx, row in enumerate(rows):
        trPr = row._tr.get_or_add_trPr()
        cant = OxmlElement("w:cantSplit")
        trPr.append(cant)
        if idx < len(rows) - 1:
            for cell in row.cells:
                for p in cell.paragraphs:
                    keep_next(p)


def drop_title_border(doc) -> None:
    """去掉 Word 内置 Title 样式自带的下方蓝色横线（本文件不需要）。"""
    for style in (doc.styles["Title"], doc.styles["Normal"]):
        try:
            pPr = style.element.get_or_add_pPr()
        except Exception:
            continue
        for bdr in pPr.findall(qn("w:pBdr")):
            pPr.remove(bdr)


def make_table(doc, headers: list[str], rows: list[list[str]], widths: list[float]) -> None:
    t = doc.add_table(rows=1, cols=len(headers))
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = False
    hdr = t.rows[0].cells
    for i, text in enumerate(headers):
        shade(hdr[i], ACCENT)
        p = hdr[i].paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_before = Pt(2)
        p.paragraph_format.space_after = Pt(2)
        set_font(p.add_run(text), 10, True, RGBColor(0xFF, 0xFF, 0xFF))
    for r, row in enumerate(rows):
        cells = t.add_row().cells
        for i, text in enumerate(row):
            if r % 2 == 1:
                shade(cells[i], "F2F6FB")
            p = cells[i].paragraphs[0]
            p.paragraph_format.space_before = Pt(2)
            p.paragraph_format.space_after = Pt(2)
            p.paragraph_format.line_spacing = 1.25
            if i == 0:
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            set_font(p.add_run(text), 9.5, i == 0, BODY_GREY)
    for row in t.rows:
        for i, w in enumerate(widths):
            row.cells[i].width = Inches(w)
    table_borders(t)
    rows_keep_together(t)
    para(doc, "", space_after=6)


def build(out: Path) -> None:
    doc = Document()
    normal = doc.styles["Normal"]
    normal.font.name = FONT
    normal.font.size = Pt(11)
    normal.element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
    for s in doc.sections:
        s.top_margin = Inches(0.9)
        s.bottom_margin = Inches(0.9)
        s.left_margin = Inches(0.95)
        s.right_margin = Inches(0.95)
    drop_title_border(doc)

    # ---------------- 封面 ----------------
    cover = OUT_DEFAULT.parent / "公告_v1.3_封面.png"
    if cover.exists():
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_after = Pt(14)
        p.add_run().add_picture(str(cover), width=Inches(6.6))

    tp = doc.add_paragraph(style="Title")
    tp.paragraph_format.space_after = Pt(2)
    tp.paragraph_format.line_spacing = 1.1
    set_font(tp.add_run("图片标签工坊 v1.3 更新公告"), 24, True, BLACK)
    tpPr = tp._p.get_or_add_pPr()
    for b in tpPr.findall(qn("w:pBdr")):
        tpPr.remove(b)

    meta = doc.add_paragraph()
    meta.paragraph_format.space_after = Pt(14)
    set_font(meta.add_run("发布日期 2026-10-05    只更新程序文件，不含模型与依赖    "
                          "github.com/HooooolyShift/imagetag-studio"), 9.5, False, META_GREY)

    para(doc, "这一版主要做两件事：把「标签体系」从一个只能看的树，改成可以直接编辑的图谱；"
              "再把积累了两个版本的标签分类、汉化、文件名规则收尾。程序依然完全在本机运行，"
              "不上传图片，也不需要联网。", 11.5, False, BODY_GREY, space_after=10)

    # ---------------- 1 图谱 ----------------
    heading(doc, "一、标签体系图谱：连线就是分类")
    para(doc, "以前的标签类型只能一个一个改，现在直接在图上连线：把标签连到「服装」节点，"
              "它的类型立刻变成服装；断开连线，类型同步撤回。")
    bullet(doc, "两种连线分色显示：灰色是分类边，青色的箭头是从属边（白裙子 → 裙子）。")
    bullet(doc, "折叠：工具栏可以全部折叠/展开，右键单个节点折叠它下方所有子节点；"
                "几个节点互连成圈也不会崩。")
    bullet(doc, "搜索定位：输入中文或英文标签回车，自动展开分类、选中并居中。")
    bullet(doc, "性能：1000 个标签、1000 条连线时，图谱打开约 0.06 秒。")
    bullet(doc, "原来单独的「标签管理」并进了图谱界面左侧，成为体系树 / 标签管理两个标签页，"
                "改类型、删标签、定位都双向联动。")
    figure(doc, SHOTS / "08_标签体系图谱.png", "图 1　标签体系图谱：连线即改分类，青色箭头表示从属关系")

    # ---------------- 2 分类与汉化 ----------------
    heading(doc, "二、标签分类与汉化收尾")
    bullet(doc, "1019 个标签全部分类完毕，落在「其它」的从 380 个降到 0。")
    bullet(doc, "新增 7 个类型：道具/物品、身体细节、文本/界面、节日/活动、性行为/互动、视角/构图、来源/杂项。")
    bullet(doc, "中文名覆盖 97%（内置词典 10275 条），界面统一显示成「中文备注（英文标签）」。")
    figure(doc, SHOTS / "06_标签管理.png", "图 2　标签管理：分类过滤 + 中文备注显示，可改类型、删除、定位到图谱")

    # ---------------- 3 从属关系 ----------------
    heading(doc, "三、标签从属关系：父与子")
    para(doc, "「更多」菜单里新增了一键生成标签从属关系，按「颜色 + 基础类」这类规则自动建关系"
              "（白裙子 → 裙子、黑裙子 → 裙子），本库一次生成 150 条，可以反复执行补全。")
    bullet(doc, "图片界面只显示最具体的标签：同时有裙子和白裙子时，只显示白裙子，不再重复刷屏。")
    bullet(doc, "筛选父标签自动带出所有子标签：筛「裙子」，白裙、黑裙、百褶裙的图片都会出现。")

    # ---------------- 4 文件名规则 ----------------
    heading(doc, "四、文件名规则调整（请留意）")
    para(doc, "这版改了标签落盘的规则，如果你一直用文件名当标签载体，建议先拿一小批试。")
    figure(doc, SHOTS / "11_设置.png", "图 3　设置 → 常规：保留原文件名、父子只留最具体、自动补回父标签三个开关",
           width_in=4.2, keep_ratio=0.56)
    bullet(doc, "默认只用标签命名：写回后是 [初音未来 泳装].png。想保留原文件名，到「设置 → 常规」"
                "勾选「写回文件名时保留原文件名」，就变回 IMG_1234 [初音未来 泳装].png。")
    bullet(doc, "父子标签只写最具体的：同时有裙子和白裙子时，文件名里只写白裙子；系列文件夹名同理。")
    bullet(doc, "库里的标签一个都不会少：检索时选「裙子」照样能找出只写了「白裙子」的图，"
                "审核结论与模型学习也照常使用全部标签。")
    bullet(doc, "重新扫描时按从属关系自动补回父标签，并且直接生效、不进待审队列；"
                "这个推断依赖库里的从属关系，换机器时记得一并导入学习包。")
    bullet(doc, "注意：首次写回之后，原文件名就找不回来了，建议先选几张试写回确认效果。")
    # ---------------- 5 分类联动 ----------------
    heading(doc, "五、分类与标签列表联动")
    para(doc, "审核台补标签、图库管理新增标签、导入、标签编辑、图谱新建标签、图谱标签管理页，"
              "全都加上了分类过滤：默认「全部（不限分类）」显示全量，选了某个分类就只列该分类下的标签"
              "（选服装，列表里就只剩围裙、泳装、女仆装这些）。")
    figure(doc, SHOTS / "05_审核台.png", "图 4　审核台：每个标签单独通过/否决，右下角「审核完毕」，分级单独一列")

    # ---------------- 6 其它改进 ----------------
    heading(doc, "六、其它改进")
    bullet(doc, "启动提醒可以勾「下次不再显示」；顺手修掉了「每次启动都弹」的根因"
                "（原来判断的是配置文件里的目录，而目录其实存在数据库里）。")
    bullet(doc, "批量删除标签：默认只删索引，可勾选同时把它们从文件名里去掉。")
    bullet(doc, "学习包导出/导入：汉化词典、类型、连线、自训练探针一起带走，跨库跨机按合并处理，不覆盖本地。")
    bullet(doc, "标签联想覆盖所有能建标签的地方，词库按图谱分类组织，选中已有标签就打到已有标签上。")
    bullet(doc, "大图预览里选中某个标签，会高亮它对应的框，方便核对偏差；框不准可以重新框选覆盖。")
    bullet(doc, "把十几个界面里重复实现的功能合并成公共实现（图片加载、缩略图、标签显示名、"
                "分类过滤、标签新建、审完标记等），以后同类问题只需要改一处。")
    figure(doc, SHOTS / "14_大图预览与框选.png", "图 5　大图预览与框选：选中标签高亮对应区域，可重新框选覆盖")

    # ---------------- 7 修复 ----------------
    heading(doc, "七、修复清单")
    bullet(doc, "图谱打不开（悬空连线导致崩溃）｜批量改分类被旧连线改回｜图谱只显示 1 个节点")
    bullet(doc, "从文件名读回的标签进了待审队列｜删掉的文件夹还留在图库列表｜库根不能删")
    bullet(doc, "一次性导入过多图片卡死｜待审数量过多时点击卡死｜标签管理页打开卡顿")
    bullet(doc, "数据库被并发写坏（新增单实例锁与并发写守卫）")

    # ---------------- 8 升级 ----------------
    h8 = doc.add_paragraph()
    h8.paragraph_format.space_before = Pt(16)
    h8.paragraph_format.space_after = Pt(6)
    set_font(h8.add_run("八、怎么更新、怎么获取"), 15.5, True, BLACK)
    h8.style = doc.styles["Heading 1"]
    for run in h8.runs:
        run.font.color.rgb = BLACK
    keep_next(h8)
    make_table(
        doc,
        ["方式", "适合谁", "怎么做", "体积"],
        [["更新包", "已经装过任意旧版本", "解压后双击 更新.cmd，填安装目录即可；只覆盖程序文件，"
                                         "不动 models / python / 数据库 / 图库", "约 5 MB"],
         ["完整安装包", "全新安装 / 换电脑离线装", "安装包里的 图片标签工坊_安装程序.exe，"
                                                 "检测硬件后自动推荐性能挡位，可自选安装路径", "4.7 GB"],
         ["源码", "想自己改 / 参与开发", "clone 仓库后跑 tools\\setup_env.ps1，"
                                         "自动建环境并从国内镜像装依赖、下模型", "约 5 GB"]],
        [1.0, 1.45, 3.15, 0.8],
    )
    para(doc, "模型来源：二次元打标用 SmilingWolf 的 WD14 系列（wd-swinv2-tagger-v3，另有 EVA02-large "
              "与 convnext v1.4 可选），自定义标签用 laion 的 CLIP ViT-B-32，人脸检测与特征用 "
              "immich-app/buffalo_l。都是公开开源模型，仓库里不重复托管，装的时候从 HuggingFace"
              "（国内走 hf-mirror 镜像）下载即可；完整安装包里的模型同样来自这些仓库，供无网络环境使用。",
         10, False, BODY_GREY, space_after=6)
    para(doc, "更新包适用于任意旧版本，升级前请先关掉正在运行的程序。", 10.5, False, BODY_GREY, space_after=10)

    # ---------------- 9 反馈 ----------------
    heading(doc, "九、问题反馈")
    para(doc, "用下来有问题或者想要新功能，欢迎在 GitHub 仓库提 issue（github.com/HooooolyShift/"
              "imagetag-studio）。标签分类、汉化名这类「数据」问题也欢迎直接反馈，可以合进下一版的学习包里。",
         11, False, BODY_GREY, space_after=4)
    para(doc, "这一版改动比较大，尤其是文件名规则，升级前记得先备份重要目录。用着不顺手的地方直接说，"
              "下个版本接着改。", 11, False, BODY_GREY, space_after=4)

    out.parent.mkdir(parents=True, exist_ok=True)
    doc.save(out)
    print("已生成:", out)


def main() -> int:
    out = Path(sys.argv[1]) if len(sys.argv) > 1 else OUT_DEFAULT
    build(out)
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
